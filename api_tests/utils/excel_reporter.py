"""
Excel Reporter for API Tests
Creates detailed Excel reports with test results
"""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from datetime import datetime
from pathlib import Path
from config.config import Config

class ExcelReporter:
    """Generate Excel reports for API tests"""
    
    def __init__(self, report_name: str):
        self.wb = Workbook()
        self.ws = self.wb.active
        self.ws.title = "API Test Results"
        self.report_name = report_name
        self.row_num = 1
        
        # Styles
        self.header_fill = PatternFill(start_color="006A4E", end_color="006A4E", fill_type="solid")
        self.header_font = Font(bold=True, color="FFFFFF", size=12)
        self.pass_fill = PatternFill(start_color="28A745", end_color="28A745", fill_type="solid")
        self.fail_fill = PatternFill(start_color="DC3545", end_color="DC3545", fill_type="solid")
        self.border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        
        self._create_header()
    
    def _create_header(self):
        """Create report header"""
        headers = ['#', 'Test Name', 'Endpoint', 'Method', 'Status Code', 
                   'Expected', 'Response Time (ms)', 'Result', 'Error Message', 'Timestamp']
        
        for col_num, header in enumerate(headers, 1):
            cell = self.ws.cell(row=1, column=col_num)
            cell.value = header
            cell.font = self.header_font
            cell.fill = self.header_fill
            cell.alignment = Alignment(horizontal='center', vertical='center')
            cell.border = self.border
        
        # Set column widths
        self.ws.column_dimensions['A'].width = 5
        self.ws.column_dimensions['B'].width = 40
        self.ws.column_dimensions['C'].width = 40
        self.ws.column_dimensions['D'].width = 10
        self.ws.column_dimensions['E'].width = 12
        self.ws.column_dimensions['F'].width = 12
        self.ws.column_dimensions['G'].width = 18
        self.ws.column_dimensions['H'].width = 12
        self.ws.column_dimensions['I'].width = 50
        self.ws.column_dimensions['J'].width = 20
        
        self.row_num = 2
    
    def add_test_result(self, test_name: str, endpoint: str, method: str,
                       status_code: int, expected_code: int, response_time: float,
                       passed: bool, error_msg: str = ""):
        """Add test result row"""
        row = [
            self.row_num - 1,
            test_name,
            endpoint,
            method,
            status_code,
            expected_code,
            round(response_time, 2),
            "PASS" if passed else "FAIL",
            error_msg,
            datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        ]
        
        for col_num, value in enumerate(row, 1):
            cell = self.ws.cell(row=self.row_num, column=col_num)
            cell.value = value
            cell.border = self.border
            cell.alignment = Alignment(horizontal='center' if col_num in [1, 4, 5, 6, 7, 8] else 'left',
                                      vertical='center')
            
            # Color code results
            if col_num == 8:  # Result column
                if passed:
                    cell.fill = self.pass_fill
                    cell.font = Font(bold=True, color="FFFFFF")
                else:
                    cell.fill = self.fail_fill
                    cell.font = Font(bold=True, color="FFFFFF")
        
        self.row_num += 1
    
    def save(self):
        """Save Excel report"""
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        filename = f"{self.report_name}_{timestamp}.xlsx"
        filepath = Config.EXCEL_REPORTS_DIR / filename
        self.wb.save(filepath)
        return filepath
