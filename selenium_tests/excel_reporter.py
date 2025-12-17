"""
Excel Report Generator for BRTA Selenium Tests
Creates detailed Excel reports with test results
"""

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from datetime import datetime
from pathlib import Path
from config import Config
import logging

logger = logging.getLogger(__name__)

class ExcelReporter:
    """Generates Excel reports for test execution"""
    
    def __init__(self, test_suite_name):
        self.test_suite_name = test_suite_name
        self.workbook = Workbook()
        self.summary_sheet = None
        self.details_sheet = None
        self.results = []
        self.initialize_workbook()  # Auto-initialize
        
    def initialize_workbook(self):
        """Initialize Excel workbook with sheets"""
        # Remove default sheet
        if 'Sheet' in self.workbook.sheetnames:
            self.workbook.remove(self.workbook['Sheet'])
        
        # Create Summary Sheet
        self.summary_sheet = self.workbook.create_sheet("Test Summary", 0)
        self._create_summary_header()
        
        # Create Details Sheet
        self.details_sheet = self.workbook.create_sheet("Test Details", 1)
        self._create_details_header()
        
        logger.info("Excel workbook initialized")
    
    def _create_summary_header(self):
        """Create header for summary sheet"""
        sheet = self.summary_sheet
        
        # Title
        sheet.merge_cells('A1:F1')
        title_cell = sheet['A1']
        title_cell.value = f"BRTA Selenium Test Report - {self.test_suite_name}"
        title_cell.font = Font(size=16, bold=True, color="FFFFFF")
        title_cell.fill = PatternFill(start_color="006A4E", end_color="006A4E", fill_type="solid")
        title_cell.alignment = Alignment(horizontal="center", vertical="center")
        sheet.row_dimensions[1].height = 30
        
        # Date/Time
        sheet.merge_cells('A2:F2')
        date_cell = sheet['A2']
        date_cell.value = f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"
        date_cell.alignment = Alignment(horizontal="center")
        
        # Statistics Headers
        stats_row = 4
        headers = ["Total Tests", "Passed", "Failed", "Skipped", "Pass Rate", "Duration"]
        for col, header in enumerate(headers, start=1):
            cell = sheet.cell(row=stats_row, column=col)
            cell.value = header
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(start_color="28A745", end_color="28A745", fill_type="solid")
            cell.alignment = Alignment(horizontal="center")
    
    def _create_details_header(self):
        """Create header for details sheet"""
        sheet = self.details_sheet
        
        # Headers
        headers = [
            "Test ID",
            "Feature Category",
            "Test Case Name",
            "Description",
            "Status",
            "Duration (s)",
            "Error Message",
            "Screenshot",
            "Timestamp"
        ]
        
        for col, header in enumerate(headers, start=1):
            cell = sheet.cell(row=1, column=col)
            cell.value = header
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(start_color="006A4E", end_color="006A4E", fill_type="solid")
            cell.alignment = Alignment(horizontal="center", vertical="center")
            
            # Set column widths
            widths = [10, 20, 30, 40, 10, 12, 50, 30, 20]
            sheet.column_dimensions[get_column_letter(col)].width = widths[col-1]
    
    def add_test_result(self, test_id, category, name, description, status, 
                       duration, error_msg="", screenshot=""):
        """Add test result to results list"""
        result = {
            'test_id': test_id,
            'category': category,
            'name': name,
            'description': description,
            'status': status,
            'duration': duration,
            'error_msg': error_msg,
            'screenshot': screenshot,
            'timestamp': datetime.now().strftime('%Y-%m-%d %H:%M:%S')
        }
        self.results.append(result)
        logger.info(f"Added test result: {name} - {status}")
    
    def write_results(self):
        """Write all results to Excel sheets"""
        # Write details
        for idx, result in enumerate(self.results, start=2):
            row = idx
            
            self.details_sheet.cell(row, 1).value = result['test_id']
            self.details_sheet.cell(row, 2).value = result['category']
            self.details_sheet.cell(row, 3).value = result['name']
            self.details_sheet.cell(row, 4).value = result['description']
            
            # Status cell with color
            status_cell = self.details_sheet.cell(row, 5)
            status_cell.value = result['status']
            status_cell.alignment = Alignment(horizontal="center")
            
            if result['status'] == 'PASS':
                status_cell.fill = PatternFill(start_color="28A745", end_color="28A745", fill_type="solid")
                status_cell.font = Font(bold=True, color="FFFFFF")
            elif result['status'] == 'FAIL':
                status_cell.fill = PatternFill(start_color="DC3545", end_color="DC3545", fill_type="solid")
                status_cell.font = Font(bold=True, color="FFFFFF")
            elif result['status'] == 'SKIPPED':
                status_cell.fill = PatternFill(start_color="6C757D", end_color="6C757D", fill_type="solid")
                status_cell.font = Font(bold=True, color="FFFFFF")
            else:
                status_cell.fill = PatternFill(start_color="FFC107", end_color="FFC107", fill_type="solid")
            
            self.details_sheet.cell(row, 6).value = round(result['duration'], 2)
            self.details_sheet.cell(row, 7).value = result['error_msg']
            # Convert Path to string for Excel compatibility
            screenshot_str = str(result['screenshot']) if result['screenshot'] else ''
            self.details_sheet.cell(row, 8).value = screenshot_str
            self.details_sheet.cell(row, 9).value = result['timestamp']
        
        # Calculate and write summary
        self._write_summary()
        
        logger.info(f"Written {len(self.results)} test results to Excel")
    
    def _write_summary(self):
        """Write summary statistics"""
        total = len(self.results)
        passed = sum(1 for r in self.results if r['status'] == 'PASS')
        failed = sum(1 for r in self.results if r['status'] == 'FAIL')
        skipped = sum(1 for r in self.results if r['status'] == 'SKIPPED')
        pass_rate = f"{(passed/total*100):.2f}%" if total > 0 else "0%"
        total_duration = sum(r['duration'] for r in self.results)
        
        # Write statistics
        stats_row = 5
        stats = [total, passed, failed, skipped, pass_rate, f"{total_duration:.2f}s"]
        
        for col, stat in enumerate(stats, start=1):
            cell = self.summary_sheet.cell(row=stats_row, column=col)
            cell.value = stat
            cell.alignment = Alignment(horizontal="center")
            cell.font = Font(size=12, bold=True)
    
    def save(self, filename=None):
        """Save Excel file"""
        # Write all results before saving
        self.write_results()
        
        if not filename:
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            filename = f"{self.test_suite_name}_{timestamp}.xlsx"
        
        filepath = Config.EXCEL_DIR / filename
        self.workbook.save(filepath)
        logger.info(f"Excel report saved: {filepath}")
        return filepath


class TestCaseTemplate:
    """Creates test case template Excel file"""
    
    @staticmethod
    def create_template():
        """Create test case template"""
        workbook = Workbook()
        
        # Remove default sheet
        if 'Sheet' in workbook.sheetnames:
            workbook.remove(workbook['Sheet'])
        
        # Create Test Cases Sheet
        sheet = workbook.create_sheet("Test Cases", 0)
        
        # Headers
        headers = [
            "Test ID",
            "Feature Category",
            "Test Case Name",
            "Test Description",
            "Preconditions",
            "Test Steps",
            "Expected Result",
            "Priority",
            "Status",
            "Notes"
        ]
        
        for col, header in enumerate(headers, start=1):
            cell = sheet.cell(row=1, column=col)
            cell.value = header
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(start_color="006A4E", end_color="006A4E", fill_type="solid")
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            
            # Set column widths
            widths = [10, 20, 30, 40, 30, 50, 30, 10, 10, 30]
            sheet.column_dimensions[get_column_letter(col)].width = widths[col-1]
        
        # Freeze first row
        sheet.freeze_panes = 'A2'
        
        # Save template
        filepath = Config.EXCEL_DIR / "test_case_template.xlsx"
        workbook.save(filepath)
        logger.info(f"Test case template created: {filepath}")
        return filepath
